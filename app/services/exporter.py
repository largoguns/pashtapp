"""Exportación a CSV y libro Excel completo (Mejora 5)."""
from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.database import Base
from app.models import Transaction, User
from app.services.balances import in_month

CSV_COLUMNS = [
    "id", "period", "date", "settlement_date", "name", "amount", "is_income", "category", "is_fixed",
    "is_settled", "installment_group_id", "installment_number", "installment_total", "notes",
]


def transactions_csv(db: Session, year: int | None, month: int | None) -> str:
    q = select(Transaction).order_by(
        Transaction.period_year, Transaction.period_month, Transaction.date, Transaction.id
    )
    if year and month:
        q = q.where(in_month(year, month))
    elif year:
        q = q.where(Transaction.period_year == year)
    buf = io.StringIO()
    buf.write("﻿")  # BOM para que Excel detecte UTF-8
    w = csv.writer(buf, delimiter=";")
    w.writerow(CSV_COLUMNS)
    for t in db.scalars(q).unique():
        w.writerow([
            t.id, f"{t.period_year:04d}-{t.period_month:02d}", t.date.isoformat(),
            t.settlement_date.isoformat() if t.settlement_date else "",
            t.name, f"{t.amount:.2f}".replace(".", ","), int(t.is_income),
            t.category.name if t.category else "", int(t.is_fixed), int(t.is_settled),
            t.installment_group_id or "", t.installment_number or "", t.installment_total or "", t.notes or "",
        ])
    return buf.getvalue()


_SKIP_TABLES = {User.__tablename__}  # Nunca exportar hashes de contraseña.
_SKIP_COLUMNS = {"key_hash"}


def full_backup_xlsx(db: Session) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0F172A")
    for table in Base.metadata.sorted_tables:
        if table.name in _SKIP_TABLES:
            continue
        cols = [c for c in table.columns if c.name not in _SKIP_COLUMNS]
        ws = wb.create_sheet(title=table.name[:31])
        ws.append([c.name for c in cols])
        for cell in ws[1]:
            cell.font, cell.fill = header_font, header_fill
        for row in db.execute(select(*cols)).all():
            ws.append(list(row))
        for i, col in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(10, min(40, len(col.name) + 4))
        ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
