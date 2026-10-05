"""Importación del libro «Pashta Saeta» (.xlsx): análisis y escritura selectiva (Mejora 1).

Lo usan ``scripts/migrate_excel.py`` (CLI) y la importación desde la web
(*Ajustes → Datos y copias*). Estructura esperada del libro (anclada en etiquetas):

  * Hojas «Enero»…«Diciembre»: cada una es un mes contable.
      - Bloque ENTRADA: «Salario» (D) → ingreso; «Resto mes anterior» (D) de enero → saldo inicial.
      - Bloque «GASTOS FIJOS …»: B casilla cobrado, C concepto, D importe.
      - Bloque «GASTOS»: B casilla, C categoría, D descripción, F importe (negativo = devolución).
  * «Gastos por Categoría»: columna A con las categorías (más la lista de validación de los meses).
  * «Gastos Electricos»: A mes, B importe.
  * «Coche», «Grueso», «Placas»: Nº de Cuota, Vencimiento, Cuota, Amortización, Intereses, Pendiente.

El libro no tiene fechas: el salario se fecha el día estimado de cobro del mes anterior (abre el
mes contable), las cuotas de préstamo en su vencimiento y el resto el día 1 de la hoja.
"""
from __future__ import annotations

import re
import unicodedata
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from statistics import median
from typing import BinaryIO

from openpyxl import load_workbook
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from app.models import (Category, Loan, LoanInstallment, MonthOpening, RecurringTemplate, Transaction,
                        UtilityReading)
from app.services.balances import initial_balance
from app.services.months import mark_opened
from app.services.parser import guess_category
from app.services.periods import estimated_start, set_start_day
from app.services.settings_store import get_setting, set_opening_balance
from app.utils import add_months, clamp_day, parse_amount

MONTHS = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
          "septiembre", "octubre", "noviembre", "diciembre"]
LOAN_SHEETS = ("Coche", "Grueso", "Placas")
CATEGORY_SHEET = "gastos por categoria"
UTILITY_SHEET = "gastos electricos"
NOTE = "[migrado Excel]"

# Categoría de cada gasto fijo del libro (el resto cae en «Fijos»).
FIXED_CATEGORY = {
    "hipoteca": "Vivienda", "seguro hogar": "Vivienda", "alarma": "Vivienda", "agua": "Vivienda",
    "telefono": "Suministros", "glp": "Suministros", "electricidad": "Luz",
    "gasolina": "Gasolina", "comida": "Comida",
    "seguro vida": "Seguros", "seguro coche": "Seguros",
    "ninos": "Familia", "extra aly": "Familia", "extra papa": "Familia",
    "impuestos": "Impuestos", "ahorro": "Ahorro",
    "residencia": "Ahorro",  # ese dinero se apartó como ahorro (Revolut)
}
CATEGORY_ICON = {
    "comida": "🛒", "ocio": "🎬", "trabajo": "💼", "farmacia": "💊", "hogar": "🏠", "regalos": "🎁",
    "ana": "🎓", "estanco": "🚬", "comida externa": "🍽️", "reintegros": "🏧", "vacaciones": "✈️",
    "devoluciones": "↩️", "fiestas": "🎉", "salario": "💶", "ahorro": "🪙", "vivienda": "🏡",
    "suministros": "🔌", "luz": "⚡", "gasolina": "⛽", "seguros": "🛡️", "familia": "🧸",
    "impuestos": "🧾", "prestamos": "🏦", "fijos": "📌", "otros": "📦",
}
FIXED_CATEGORIES = {"vivienda", "suministros", "luz", "seguros", "familia", "impuestos", "prestamos", "fijos", "salario"}
PALETTE = ["#3987e5", "#d95926", "#199e70", "#c98500", "#d55181", "#008300", "#9085e9", "#e66767"]


def norm(v) -> str:
    s = unicodedata.normalize("NFKD", str(v or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().lower()


def num(v) -> float | None:
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return parse_amount(str(v)) if re.search(r"\d", str(v)) else None


def as_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    return v if isinstance(v, date) else None


# ----------------------------------------------------------------------------- estructuras

@dataclass
class Tx:
    year: int
    month: int
    kind: str              # income | fixed | variable
    name: str
    amount: float          # con signo: + ingreso / - gasto (devolución: + no ingreso)
    settled: bool
    category: str
    notes: str = NOTE


@dataclass
class LoanData:
    name: str
    rows: list[dict]
    capital: float
    rate: float
    term: int
    fee: float
    start: date


@dataclass
class Result:
    year: int
    categories: list[str] = field(default_factory=list)
    transactions: list[Tx] = field(default_factory=list)
    loans: list[LoanData] = field(default_factory=list)
    utility: dict[int, float] = field(default_factory=dict)
    opening: float | None = None
    carry: dict[int, float] = field(default_factory=dict)  # «Resto mes anterior» de cada mes (validación)
    months: list[int] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def find_label(ws, pattern: str, col: int | None = None) -> int | None:
    """Fila (1-based) de la primera celda cuyo texto normalizado casa con ``pattern``."""
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and (col is None or c.column == col) and re.search(pattern, norm(c.value)):
                return c.row
    return None


# ----------------------------------------------------------------------------- hojas

def parse_month(ws, year: int, month: int, today: date, res: Result) -> None:
    v = lambda r, c: ws.cell(r, c).value  # noqa: E731
    is_past_or_current = (year, month) <= (today.year, today.month)

    r_salary = find_label(ws, r"^salario", col=3)
    r_carry = find_label(ws, r"^resto mes anterior", col=3)
    r_fixed = find_label(ws, r"^gastos fijos", col=3)
    r_var = find_label(ws, r"^gastos$", col=3)
    if not (r_fixed and r_var):
        res.warnings.append(f"[{ws.title}] no se encuentran los bloques GASTOS FIJOS / GASTOS")
        return

    if r_salary and num(v(r_salary, 4)):
        # El salario no tiene casilla: se da por cobrado en meses pasados y en el actual.
        res.transactions.append(Tx(year, month, "income", "Salario", abs(num(v(r_salary, 4))),
                                   is_past_or_current, "Salario"))
    if r_carry and num(v(r_carry, 4)) is not None:
        res.carry[month] = num(v(r_carry, 4))

    for r in range(r_fixed + 1, r_var):
        name, amount = v(r, 3), num(v(r, 4))
        if not isinstance(name, str) or not name.strip() or not amount:
            continue
        name = name.strip()
        res.transactions.append(Tx(year, month, "fixed", name, -amount, v(r, 2) is True,
                                   FIXED_CATEGORY.get(norm(name), "Fijos")))

    empty = 0
    for r in range(r_var + 1, ws.max_row + 1):
        cat, desc, amount = v(r, 3), v(r, 4), num(v(r, 6))
        if amount in (None, 0):
            empty = empty + 1 if cat is None and desc is None else 0
            if empty > 40:
                break
            continue
        empty = 0
        cat = str(cat).strip() if cat else None
        desc = str(desc).strip() if desc else None
        if not cat:
            cat = guess_category(desc or "")
        res.transactions.append(Tx(year, month, "variable", desc or cat, -amount, v(r, 2) is True, cat))
    res.months.append(month)


def parse_categories(wb, res: Result) -> None:
    names: list[str] = []
    for ws in wb.worksheets:
        if norm(ws.title) == CATEGORY_SHEET:
            for (val,) in ws.iter_rows(min_row=2, max_col=1, values_only=True):
                if isinstance(val, str) and val.strip():
                    names.append(val.strip())
    # Lista desplegable (validación de datos) de la columna de categoría de los meses.
    for ws in wb.worksheets:
        if norm(ws.title) in MONTHS and ws.data_validations:
            for dv in ws.data_validations.dataValidation:
                if dv.type == "list" and dv.formula1 and dv.formula1.startswith('"'):
                    names.extend(x.strip() for x in dv.formula1.strip('"').split(",") if x.strip())
    seen = set()
    res.categories = [n for n in names if not (norm(n) in seen or seen.add(norm(n)))]


def parse_utility(wb, res: Result, settled_electricity: dict[int, bool]) -> None:
    ws = next((w for w in wb.worksheets if norm(w.title) == UTILITY_SHEET), None)
    if ws is None:
        res.warnings.append("No hay hoja «Gastos Electricos»")
        return
    for name, amount in ws.iter_rows(min_col=1, max_col=2, values_only=True):
        if norm(name) in MONTHS and num(amount):
            m = MONTHS.index(norm(name)) + 1
            # Los meses sin el cargo «Electricidad» conciliado son estimaciones del libro: se omiten.
            if settled_electricity.get(m, False):
                res.utility[m] = round(num(amount), 2)


def parse_loan(ws, res: Result) -> None:
    header = None
    for row in ws.iter_rows(max_row=10):
        heads = [norm(c.value) for c in row]
        if any("interes" in h for h in heads) and any("pendiente" in h for h in heads):
            header = row[0].row
            cols = {
                "number": next(i for i, h in enumerate(heads) if h.startswith("n")),
                "due": next(i for i, h in enumerate(heads) if "vencim" in h or "fecha" in h),
                "payment": next(i for i, h in enumerate(heads) if "cuota" in h and not h.startswith("n")),
                "capital": next(i for i, h in enumerate(heads) if "amortiz" in h),
                "interest": next(i for i, h in enumerate(heads) if "interes" in h),
                "remaining": next(i for i, h in enumerate(heads) if "pendiente" in h),
            }
            break
    if header is None:
        res.warnings.append(f"[{ws.title}] sin cuadro de amortización reconocible")
        return
    rows = []
    for vals in ws.iter_rows(min_row=header + 1, values_only=True):
        payment, due = num(vals[cols["payment"]]), as_date(vals[cols["due"]])
        if payment is None or due is None:
            continue
        rows.append({
            "number": int(num(vals[cols["number"]]) or len(rows) + 1), "due": due, "payment": payment,
            "capital": num(vals[cols["capital"]]) or 0.0, "interest": num(vals[cols["interest"]]) or 0.0,
            "remaining": num(vals[cols["remaining"]]) or 0.0,
        })
    if not rows:
        res.warnings.append(f"[{ws.title}] cuadro vacío")
        return
    # TIN: mediana de interés / capital vivo de las primeras cuotas (la 1ª puede ser un periodo parcial).
    rates, prev = [], rows[0]["remaining"] + rows[0]["capital"]
    for r in rows[:12]:
        if prev > 0 and r["interest"] > 0:
            rates.append(r["interest"] / prev * 1200)
        prev = r["remaining"]
    res.loans.append(LoanData(
        name=ws.title.strip(),
        rows=rows,
        capital=round(rows[0]["remaining"] + rows[0]["capital"], 2),
        rate=round(median(rates), 3) if rates else 0.0,
        term=len(rows),
        fee=Counter(round(r["payment"], 2) for r in rows).most_common(1)[0][0],
        start=add_months(rows[0]["due"], -1),
    ))


def analyse(source: str | Path | BinaryIO, year: int, today: date) -> Result:
    """Lee el libro (ruta o fichero abierto) y devuelve lo que se importaría, sin tocar la BD."""
    wb = load_workbook(source, data_only=True)  # valores calculados de las fórmulas
    res = Result(year=year)
    parse_categories(wb, res)
    for ws in wb.worksheets:
        if norm(ws.title) in MONTHS:
            parse_month(ws, year, MONTHS.index(norm(ws.title)) + 1, today, res)
        elif ws.title.strip() in LOAN_SHEETS:
            parse_loan(ws, res)
    elec = {t.month: t.settled for t in res.transactions if t.kind == "fixed" and norm(t.name) == "electricidad"}
    parse_utility(wb, res, elec)
    if 1 in res.carry:
        res.opening = res.carry[1]
    return res


def is_loan_payment(name: str, loans: list[LoanData]) -> LoanData | None:
    """«Coche», «Préstamo Grueso»… (pero no «Seguro Coche»)."""
    n = norm(name)
    for loan in loans:
        if re.fullmatch(rf"(prestamo\s+)?{re.escape(norm(loan.name))}", n):
            return loan
    return None




# ----------------------------------------------------------------------------- escritura


@dataclass
class ImportOptions:
    today: date
    months: set[int] | None = None  # meses (hojas) a importar; None = todos
    on_existing: str = "skip"       # mes que ya tiene movimientos: "skip" | "replace"
    loans: bool = True
    replace_loans: bool = False     # préstamo con el mismo nombre: reemplazar su cuadro
    utility: bool = True
    templates: bool = True          # crear plantillas de fijos si no hay ninguna
    start_day: int = 27
    reset: bool = False             # vaciar categorías, movimientos, préstamos, plantillas y luz


@dataclass
class ImportReport:
    year: int
    imported: list[int] = field(default_factory=list)
    replaced: list[int] = field(default_factory=list)
    skipped: list[int] = field(default_factory=list)
    transactions: int = 0
    loans_created: list[str] = field(default_factory=list)
    loans_replaced: list[str] = field(default_factory=list)
    loans_kept: list[str] = field(default_factory=list)
    utility: int = 0
    templates: int = 0
    first_import: bool = False
    checks: dict[int, tuple[float, float]] = field(default_factory=dict)


def existing_counts(db: Session, year: int) -> dict[int, int]:
    """Movimientos que ya hay en cada mes contable del año."""
    rows = db.execute(
        select(Transaction.period_month, func.count(Transaction.id))
        .where(Transaction.period_year == year).group_by(Transaction.period_month)
    ).all()
    return dict(rows)


def month_overview(res: Result) -> dict[int, dict]:
    """Resumen por hoja para la vista previa."""
    out: dict[int, dict] = {m: {"income": 0.0, "fixed": 0.0, "variable": 0.0, "count": 0, "settled": 0}
                            for m in res.months}
    for t in res.transactions:
        row = out[t.month]
        row["count"] += 1
        row["settled"] += t.settled
        row[t.kind] += abs(t.amount) if t.kind == "income" else -t.amount
    return out


def import_workbook(db: Session, res: Result, opts: ImportOptions) -> ImportReport:
    report = ImportReport(year=res.year)
    if opts.reset:
        for model in (Transaction, LoanInstallment, Loan, RecurringTemplate, UtilityReading, MonthOpening, Category):
            db.execute(delete(model))
        db.flush()
    report.first_import = not db.scalar(select(func.count(Transaction.id)))

    cats: dict[str, Category] = {norm(c.name): c for c in db.scalars(select(Category))}

    def category(name: str) -> Category:
        key = norm(name)
        if key not in cats:
            cats[key] = Category(
                name=name, icon=CATEGORY_ICON.get(key, "tag"), color_hex=PALETTE[len(cats) % len(PALETTE)],
                is_fixed_default=key in FIXED_CATEGORIES,
            )
            db.add(cats[key])
            db.flush()
        return cats[key]

    for name in res.categories:
        category(name)

    # --- Préstamos: los existentes se conservan salvo que se pida reemplazarlos.
    installments: dict[tuple[str, int, int], LoanInstallment] = {}
    existing_loans = {norm(ln.name): ln for ln in db.scalars(select(Loan)).unique()}
    for ln in res.loans:
        key = norm(ln.name)
        current = existing_loans.get(key)
        if current is not None and (not opts.loans or not opts.replace_loans):
            report.loans_kept.append(current.name)
            for inst in current.installments:
                installments[(key, inst.due_date.year, inst.due_date.month)] = inst
            continue
        if not opts.loans:
            continue
        if current is not None:
            ids = [i.id for i in current.installments]
            if ids:
                for tx in db.scalars(select(Transaction).where(Transaction.loan_installment_id.in_(ids))):
                    tx.loan_installment_id = None
            db.delete(current)
            db.flush()
            report.loans_replaced.append(ln.name)
        else:
            report.loans_created.append(ln.name)
        loan = Loan(name=ln.name, initial_capital=ln.capital, annual_interest_rate=ln.rate,
                    term_months=ln.term, start_date=ln.start, monthly_fee=ln.fee, active=True)
        for r in ln.rows:
            inst = LoanInstallment(
                installment_number=r["number"], due_date=r["due"], payment_amount=r["payment"],
                capital_amount=r["capital"], interest_amount=r["interest"], remaining_capital=r["remaining"],
                is_paid=r["due"] <= opts.today,
            )
            loan.installments.append(inst)
            installments[(key, r["due"].year, r["due"].month)] = inst
        loan.active = any(not i.is_paid for i in loan.installments)
        db.add(loan)
    db.flush()

    # --- Meses: cada hoja es un mes contable.
    counts = existing_counts(db, res.year)
    selected = [m for m in res.months if opts.months is None or m in opts.months]
    for m in selected:
        if counts.get(m):
            if opts.on_existing != "replace":
                report.skipped.append(m)
                continue
            db.execute(delete(Transaction).where(Transaction.period_year == res.year, Transaction.period_month == m))
            report.replaced.append(m)
        else:
            report.imported.append(m)

    to_write = set(report.imported) | set(report.replaced)
    loan_cat = category("Préstamos")
    loans_due_day = {norm(ln.name): ln.rows[0]["due"].day for ln in res.loans}
    for t in res.transactions:
        if t.month not in to_write:
            continue
        loan = is_loan_payment(t.name, res.loans) if t.kind == "fixed" else None
        day = loans_due_day.get(norm(loan.name), 1) if loan else 1
        d = estimated_start(t.year, t.month, opts.start_day) if t.kind == "income" else clamp_day(t.year, t.month, day)
        cat = loan_cat if loan else category(t.category)
        tx = Transaction(
            period_year=t.year, period_month=t.month,  # la hoja del Excel es el mes contable
            date=d, settlement_date=d, name=t.name, amount=round(t.amount, 2), is_income=t.kind == "income",
            category_id=cat.id, is_fixed=t.kind == "fixed", is_settled=t.settled, notes=t.notes,
            savings_account_id=None if t.kind == "income" else cat.savings_account_id,
        )
        if loan:
            inst = installments.get((norm(loan.name), t.year, t.month))
            if inst is not None:
                db.flush()
                tx.loan_installment_id = inst.id
                tx.notes = f"[Cuota {inst.installment_number}/{len(loan.rows)}] {NOTE}"
                inst.is_paid = t.settled  # el libro manda en los meses importados
        db.add(tx)
        report.transactions += 1
    for m in to_write:
        mark_opened(db, res.year, m)  # sus fijos ya están importados: no regenerar

    # --- Luz.
    if opts.utility:
        for m, amount in res.utility.items():
            row = db.scalar(select(UtilityReading).where(UtilityReading.year == res.year, UtilityReading.month == m))
            row = row or UtilityReading(year=res.year, month=m)
            row.amount, row.notes = amount, row.notes or NOTE
            db.add(row)
            report.utility += 1

    # --- Ajustes iniciales: sólo en la primera importación (no pisar la configuración existente).
    if report.first_import and to_write:
        set_opening_balance(db, res.opening if 1 in to_write and res.opening is not None else 0.0,
                            res.year, min(to_write))
    if report.first_import or get_setting(db, "period_start_day") is None:
        set_start_day(db, opts.start_day)

    # --- Plantillas de fijos a partir del último mes importado (si no hay ninguna).
    if opts.templates and to_write and not db.scalar(select(func.count(RecurringTemplate.id))):
        last = max(to_write)
        for t in res.transactions:
            if t.month != last or t.kind == "variable" or is_loan_payment(t.name, res.loans):
                continue
            db.add(RecurringTemplate(name=t.name, default_amount=abs(t.amount), category_id=category(t.category).id,
                                     day_of_month=opts.start_day if t.kind == "income" else 1,
                                     is_income=t.kind == "income", active=True))
            report.templates += 1
    db.flush()

    # Validación: arrastre calculado frente al «Resto mes anterior» del libro.
    for m in sorted(to_write):
        if m in res.carry and m != min(res.months):
            report.checks[m] = (initial_balance(db, res.year, m), res.carry[m])
    return report
