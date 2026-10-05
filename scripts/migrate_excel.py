#!/usr/bin/env python3
"""Importa el libro «Pashta Saeta» (.xlsx) en PashtAPP desde la línea de comandos (Mejora 1).

También se puede hacer desde la web: *Ajustes → Datos y copias → Importar desde Excel*.

Uso:
    python scripts/migrate_excel.py "Pashta Saeta 2026.xlsx" --dry-run        # analizar sin escribir
    python scripts/migrate_excel.py "Pashta Saeta 2026.xlsx"                  # importar (omite meses con datos)
    python scripts/migrate_excel.py libro.xlsx --months 10,11 --replace       # reimportar sólo esas hojas
    python scripts/migrate_excel.py libro.xlsx --reset                        # vaciar datos e importar todo
    python scripts/migrate_excel.py libro.xlsx --inspect                      # volcar celdas de cada hoja

La estructura esperada del libro está descrita en ``app/services/excel_import.py``.
"""
from __future__ import annotations

import argparse
import os
import re
import sys
from collections import Counter
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook  # noqa: E402

MONTHS = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
          "septiembre", "octubre", "noviembre", "diciembre"]


def summary(res) -> None:
    kinds = Counter(t.kind for t in res.transactions)
    print(f"Año {res.year} · meses {res.months}")
    print(f"Categorías: {', '.join(res.categories)}")
    print(f"Movimientos: {len(res.transactions)} (ingresos {kinds['income']}, fijos {kinds['fixed']}, "
          f"variables {kinds['variable']}; conciliados {sum(t.settled for t in res.transactions)})")
    for ln in res.loans:
        print(f"Préstamo {ln.name}: {ln.capital:.2f} € · {ln.rate} % TIN · {ln.term} cuotas de {ln.fee:.2f} € "
              f"· {ln.rows[0]['due']} → {ln.rows[-1]['due']}")
    print(f"Luz: {len(res.utility)} meses · saldo inicial (Resto mes anterior de enero): {res.opening}")
    for w in res.warnings:
        print("AVISO:", w)


def inspect(path: Path, max_rows: int = 40) -> None:
    wb = load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        print(f"\n=== «{ws.title}» ({ws.max_row}×{ws.max_column}) ===")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                break
            cells = [f"{chr(65 + c)}:{v!r}" for c, v in enumerate(row) if v is not None]
            if cells:
                print(f"{i + 1:>4} │ " + "  ".join(cells)[:240])


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", type=Path, nargs="?", default=Path("Pashta Saeta 2026.xlsx"))
    ap.add_argument("--year", type=int, help="Año de las hojas mensuales (por defecto, el del nombre del fichero)")
    ap.add_argument("--db", help="DATABASE_URL de destino (por defecto, la del entorno)")
    ap.add_argument("--today", type=date.fromisoformat, help="Fecha de referencia (AAAA-MM-DD) para cuotas pagadas")
    ap.add_argument("--months", help="Hojas a importar, p. ej. 10,11,12 (por defecto, todas)")
    ap.add_argument("--replace", action="store_true", help="Reemplazar los meses que ya tengan movimientos")
    ap.add_argument("--replace-loans", action="store_true", help="Reemplazar préstamos ya existentes")
    ap.add_argument("--inspect", action="store_true", help="Vuelca las primeras filas de cada hoja y sale")
    ap.add_argument("--dry-run", action="store_true", help="Analiza y muestra el resumen sin escribir")
    ap.add_argument("--reset", action="store_true", help="Vacía categorías, movimientos, préstamos, plantillas y luz")
    ap.add_argument("--no-templates", action="store_true", help="No crear plantillas de fijos")
    ap.add_argument("--start-day", type=int, default=27,
                    help="Día estimado de cobro del salario, que abre el mes contable (1 = meses naturales)")
    args = ap.parse_args()

    if not args.xlsx.exists():
        sys.exit(f"No existe el fichero: {args.xlsx}")
    if args.inspect:
        inspect(args.xlsx)
        return
    if args.db:
        os.environ["DATABASE_URL"] = args.db

    from app.database import SessionLocal, init_db
    from app.services.excel_import import ImportOptions, analyse, import_workbook
    from app.utils import today as today_fn

    today = args.today or today_fn()
    year = args.year or int((re.search(r"(20\d\d)", args.xlsx.name) or [None, today.year])[1])
    res = analyse(args.xlsx, year, today)
    summary(res)
    if args.dry_run:
        return

    months = {int(m) for m in args.months.split(",")} if args.months else None
    opts = ImportOptions(
        today=today, months=months, on_existing="replace" if args.replace else "skip",
        replace_loans=args.replace_loans, templates=not args.no_templates,
        start_day=args.start_day, reset=args.reset,
    )
    init_db()
    with SessionLocal() as db:
        report = import_workbook(db, res, opts)
        db.commit()

    name = lambda ms: ", ".join(MONTHS[m - 1] for m in ms) or "—"  # noqa: E731
    print(f"\n✔ Importación completada: {report.transactions} movimientos.")
    print(f"  Meses nuevos: {name(report.imported)} · reemplazados: {name(report.replaced)} "
          f"· omitidos (ya tenían datos): {name(report.skipped)}")
    print(f"  Préstamos creados: {report.loans_created or '—'} · reemplazados: {report.loans_replaced or '—'} "
          f"· conservados: {report.loans_kept or '—'}")
    if report.checks:
        print("\nValidación del arrastre de saldo (PashtAPP vs «Resto mes anterior» del libro):")
        for m, (ours, theirs) in report.checks.items():
            flag = "✓" if abs(ours - theirs) < 0.01 else f"Δ {ours - theirs:+.2f}"
            print(f"  {MONTHS[m - 1]:<11} {ours:>10.2f} {theirs:>10.2f}  {flag}")


if __name__ == "__main__":
    main()
