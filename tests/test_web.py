from datetime import date

from app.models import Category, Loan, Transaction
from app.services.auth import create_api_key
from app.services.loans import create_loan


def test_login_required(anon):
    r = anon.get("/", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/login"
    r = anon.patch("/transactions/1/toggle-settled", headers={"HX-Request": "true"})
    assert r.status_code == 401 and r.headers["HX-Redirect"] == "/login"
    assert anon.post("/login", data={"username": "admin", "password": "mal"}).status_code == 401


def test_pages_render(client):
    for url in ["/", "/?y=2026&m=2", "/loans", "/settings", "/settings?tab=categories", "/settings?tab=templates",
                "/settings?tab=utility", "/settings?tab=api", "/settings?tab=data", "/manifest.json", "/sw.js", "/healthz"]:
        assert client.get(url).status_code == 200, url


def test_quick_add_and_toggle(client, db):
    cat = db.query(Category).filter_by(name="Comida").one()
    r = client.post("/transactions", data={"amount": "12,50", "category_id": str(cat.id), "mode": "pending",
                                           "view_y": "2026", "view_m": "10", "op_date": "2026-10-05"})
    assert r.status_code == 200 and 'id="mobile-header"' in r.text and "hx-swap-oob" in r.text
    tx = db.query(Transaction).one()
    assert (tx.amount, tx.name, tx.is_settled) == (-12.5, "Comida", False)
    r = client.patch(f"/transactions/{tx.id}/toggle-settled?view_y=2026&view_m=10")
    assert r.status_code == 200
    db.expire_all()
    assert db.get(Transaction, tx.id).is_settled


def test_quick_expense_api(client, db):
    token = create_api_key(db, "test")
    r = client.post("/api/v1/quick-expense", json={"text": "Pago con tarjeta de 42,30 EUR en REPSOL..."},
                    headers={"X-API-Key": token})
    assert r.status_code == 201, r.text
    assert r.json()["amount"] == -42.3 and r.json()["category"] == "Gasolina"
    r = client.post("/api/v1/quick-expense", headers={"X-API-Key": token},
                    json={"amount": 14.50, "concept": "Mercadona", "category": "Comida", "is_settled": False})
    assert r.status_code == 201 and r.json()["category"] == "Comida"
    assert client.post("/api/v1/quick-expense", json={"amount": 1}, headers={"X-API-Key": "nope"}).status_code == 401
    assert client.post("/api/v1/quick-expense", json={"text": "hola"}, headers={"X-API-Key": token}).status_code == 422


def test_simulator_endpoint(client, db):
    loan = create_loan(db, name="Placas", initial_capital=5500, annual_interest_rate=6.45, term_months=84,
                       start_date=date(2022, 12, 5), mark_paid_until=date(2026, 10, 5))
    db.commit()
    r = client.post(f"/api/v1/loans/{loan.id}/simulate-prepayment",
                    json={"prepayment_amount": 1000, "strategy": "REDUCE_FEE"})
    assert r.status_code == 200 and r.json()["interest_savings"] > 0
    r = client.post(f"/api/v1/loans/{loan.id}/simulate-prepayment", headers={"HX-Request": "true"},
                    data={"prepayment_amount": "1000", "strategy": "REDUCE_TERM"})
    assert "Ahorro neto de intereses" in r.text
    assert client.post(f"/api/v1/loans/{loan.id}/simulate-prepayment",
                       json={"prepayment_amount": 99999}).status_code == 422
    assert db.query(Loan).count() == 1


def test_exports_and_backup(client, db):
    from openpyxl import load_workbook
    import io

    db.add(Transaction(date=date(2026, 3, 1), name="Algo", amount=-5)); db.commit()
    r = client.get("/export/transactions.csv?year=2026&month=3")
    assert r.status_code == 200 and "Algo" in r.text and "-5,00" in r.text
    r = client.get("/export/full-backup.xlsx")
    wb = load_workbook(io.BytesIO(r.content))
    assert "transactions" in wb.sheetnames and "users" not in wb.sheetnames
    r = client.post("/backups/run", follow_redirects=False)
    assert r.status_code == 303
    assert "pashtapp_" in client.get("/settings?tab=data").text


def test_password_hash_accepts_base64(monkeypatch):
    import base64

    import pytest

    from app.config import _decode_password_hash
    from app.services.auth import hash_password, verify_password

    h = hash_password("clave")
    b64 = base64.b64encode(h.encode()).decode()
    assert "$" not in b64
    assert _decode_password_hash(b64) == h and verify_password("clave", _decode_password_hash(b64))
    assert _decode_password_hash(h) == h
    assert _decode_password_hash(f"'{h}'") == h
    assert _decode_password_hash(None) is None
    # Un valor que no es hash ni base64 de hash falla al arrancar, no en el login.
    with pytest.raises(RuntimeError):
        _decode_password_hash("no-es-un-hash")
