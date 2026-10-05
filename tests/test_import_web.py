"""Importación desde la web con un libro sintético con el formato de Pashta Saeta."""
import io
import re
from datetime import datetime

from openpyxl import Workbook

from app.models import Loan, Transaction, UtilityReading


def make_book(salary=2000.0, food=40.0) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, fixed_settled in (("Enero", True), ("Febrero", False)):
        ws = wb.create_sheet(name)
        ws["C2"] = "ENTRADA"
        ws["C3"], ws["D3"] = "Salario", salary
        ws["C4"], ws["D4"] = "Resto mes anterior", 100.0
        ws["C7"] = "GASTOS FIJOS VIVIENDA"
        ws["B8"], ws["C8"], ws["D8"] = fixed_settled, "Hipoteca", 400.0
        ws["B9"], ws["C9"], ws["D9"] = True, "Préstamo Coche", 250.0
        # La luz del mes sólo se importa si su cargo «Electricidad» está conciliado.
        ws["B10"], ws["C10"], ws["D10"] = True, "Electricidad", 60.0
        ws["C25"] = "GASTOS"
        ws["B27"], ws["C27"], ws["D27"], ws["F27"] = True, "Comida", "Súper", food
        ws["B28"], ws["C28"], ws["D28"], ws["F28"] = True, "Devoluciones", "Tienda", -10.0
    el = wb.create_sheet("Gastos Electricos")
    el.append(["Enero", 60.0])
    el.append(["Febrero", 55.0])
    ws = wb.create_sheet("Coche")
    ws.append(["Nº de Cuota", "Vencimiento", "Cuota Mensual", "Amortización Capital", "Intereses", "Capital Pendiente"])
    remaining = 1000.0
    for i in range(1, 5):
        interest = round(remaining * 0.005, 2)
        capital = round(250 - interest, 2) if i < 4 else remaining
        remaining = round(remaining - capital, 2)
        ws.append([i, datetime(2026, i, 5), round(capital + interest, 2), capital, interest, remaining])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(client, data: bytes, name="Pashta Saeta 2026.xlsx"):
    r = client.post("/import", files={"file": (name, data, "application/octet-stream")}, follow_redirects=False)
    assert r.status_code == 303, r.text
    return re.search(r"/import/([0-9a-f]{32})", r.headers["location"]).group(1)


def test_upload_preview_and_import(client, db):
    token = upload(client, make_book())
    r = client.get(f"/import/{token}?year=2026")
    assert r.status_code == 200
    assert "Enero" in r.text and "Febrero" in r.text and "Coche" in r.text and "Vacío" in r.text

    r = client.post(f"/import/{token}", data={"year": "2026", "months": ["1", "2"], "loans": "true",
                                              "utility": "true", "templates": "true", "start_day": "27"})
    assert r.status_code == 200 and "Importación completada" in r.text
    assert db.query(Transaction).filter_by(period_year=2026).count() == 12
    salary = db.query(Transaction).filter_by(name="Salario", period_month=2).one()
    assert salary.date.isoformat() == "2026-01-27" and salary.is_income
    loan_tx = db.query(Transaction).filter_by(name="Préstamo Coche", period_month=1).one()
    assert loan_tx.loan_installment_id is not None
    assert db.query(Loan).count() == 1 and db.query(UtilityReading).count() == 2
    # El fichero temporal se borra tras importar.
    assert client.get(f"/import/{token}?year=2026").status_code == 404
    # Se hizo copia de seguridad antes de importar.
    assert "pashtapp_" in client.get("/settings?tab=data").text


def test_existing_months_skip_or_replace(client, db):
    token = upload(client, make_book())
    client.post(f"/import/{token}", data={"year": "2026", "months": ["1", "2"], "loans": "true"})
    # Segundo libro con otros importes: por defecto se omiten los meses con datos…
    token = upload(client, make_book(food=99.0))
    r = client.get(f"/import/{token}?year=2026")
    assert "Ya tiene 6" in r.text
    client.post(f"/import/{token}", data={"year": "2026", "months": ["1", "2"], "loans": "true",
                                          "on_existing": "skip"})
    assert db.query(Transaction).filter_by(name="Súper").first().amount == -40
    # …y con «reemplazar» sólo cambia el mes elegido, sin duplicar préstamos.
    token = upload(client, make_book(food=99.0))
    r = client.post(f"/import/{token}", data={"year": "2026", "months": ["2"], "loans": "true",
                                              "on_existing": "replace"})
    assert "Reemplazados" in r.text
    db.expire_all()
    amounts = {t.period_month: t.amount for t in db.query(Transaction).filter_by(name="Súper")}
    assert amounts == {1: -40, 2: -99}
    assert db.query(Transaction).count() == 12 and db.query(Loan).count() == 1


def test_rejects_invalid_files(client):
    r = client.post("/import", files={"file": ("notas.txt", b"hola", "text/plain")})
    assert r.status_code == 422 and ".xlsx" in r.text
    r = client.post("/import", files={"file": ("roto.xlsx", b"esto no es un zip", "application/octet-stream")})
    assert r.status_code == 422 and "No se ha podido leer" in r.text
    assert client.get("/import/../../etc/passwd?year=2026").status_code == 404
    assert client.get("/import/" + "0" * 32 + "?year=2026").status_code == 404
